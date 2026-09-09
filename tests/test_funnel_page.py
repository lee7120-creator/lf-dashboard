"""「02. 첫구매 퍼널별 상세 실적」 렌더·계산 테스트.

이 화면은 **비율 두 칸의 출처가 다르다**. 가입율은 두 카운트에서 계산하고, 당일가입CR은
파일 값을 그대로 쓴다(분자인 당일가입 첫구매 고객수가 데이터에 없다). 둘을 한쪽으로
통일해 버리면 화면 숫자가 통째로 바뀌는데, 증상이 '값이 좀 다르네'로만 보여서 스모크로는
절대 안 잡힌다. 그래서 파일 값과 계산 값을 **일부러 다르게** 심어 두고 어느 쪽이 나오는지
본다.

로컬 실행:
    python tests/test_funnel_page.py
"""
import datetime
import inspect
import io
import json
import os
import pathlib
import re
import shutil
import sys
import tempfile

import numpy as np
import pandas as pd

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import weekly_report as W                        # noqa: E402

APP = ROOT / "weekly_report.py"
TIMEOUT = 300
PAGE = "02. 첫구매 퍼널별 상세 실적"

STORE_COLS = ["gran", "metric", "segment", "year", "label", "close", "sortkey", "value"]
YEARS = (2025, 2026)
MONTHS = range(1, 10)
WEEKS = range(1, 5)

# 연도별 원값 — 채널 두 개(직접·광고)가 전체와 다른 비율을 갖도록 일부러 갈라 놨다.
#   전체 가입율 2026 = 1000/100000 = 1.00% · 직접 = 800/40000 = 2.00% · 광고 = 200/60000 = 0.33%
# 파일이 주는 「가입율」 지표는 50%로 심어 둔다 — 화면이 파일 값을 쓰면 대번에 드러난다.
VALS = {
    2026: {"*TOTAL": dict(traffic=100_000, join=1_000, cust=500, rev=50_000_000),
           "직접":   dict(traffic=40_000, join=800, cust=300, rev=30_000_000),
           "광고":   dict(traffic=60_000, join=200, cust=200, rev=20_000_000)},
    2025: {"*TOTAL": dict(traffic=80_000, join=500, cust=400, rev=32_000_000),
           "직접":   dict(traffic=30_000, join=400, cust=250, rev=20_000_000),
           "광고":   dict(traffic=50_000, join=100, cust=150, rev=12_000_000)},
}
FILE_RATE = 0.5           # 파일이 주는 가입율 — 계산값(1.00%)과 확연히 다르게
DAILY_CR = 0.0725         # 당일가입CR — 역산이 불가능한 값이라 파일 값이 그대로 나와야
FILE_AOV = 111_111        # 파일이 주는 객단가 — 거래액/고객수(100,000원)와 다르게
PUSH_DAILY = {1: 100.0, 8: 200.0, 15: 300.0, 22: 400.0}    # 일평균 250


def _rows_for(gran, label, sortkey, year):
    out = []

    def add(met, seg, v):
        out.append(dict(gran=gran, metric=met, segment=seg, year=year, label=label,
                        close="final", sortkey=sortkey, value=float(v)))
    for seg, d in VALS[year].items():
        add("비회원트래픽", seg, d["traffic"])
        add("가입자수", seg, d["join"])
        add("첫구매 고객수", seg, d["cust"])
        add("첫구매 거래액", seg, d["rev"])
        add("가입율", seg, FILE_RATE)
        add("당일가입CR", seg, DAILY_CR)
        add("첫구매 객단가", seg, FILE_AOV)
    return out


def synth_store(with_push=True):
    rows = []
    for y in YEARS:
        for mo in MONTHS:
            rows += _rows_for("월", f"{mo}월", y * 10000 + mo * 100, y)
            for wk in WEEKS:
                rows += _rows_for("주", f"{mo:02d}월 {wk}주차",
                                  y * 10000 + mo * 100 + wk, y)
            if with_push:
                # 앱푸시 수신동의는 원천이 일자 헤더 표라 **일별로만** 쌓인다
                for dd, v in PUSH_DAILY.items():
                    rows.append(dict(gran="일", metric="앱푸시수신동의", segment="*TOTAL",
                                     year=y, label=f"{mo}/{dd}", close="final",
                                     sortkey=y * 10000 + mo * 100 + dd, value=v))
    return pd.DataFrame(rows)[STORE_COLS]


# 「가방」은 **두 조직이 나눠 갖는다** — 조직 합산 경로를 실제로 밟게 하려고 그렇다.
# 한 조직에만 있는 카테고리뿐이면 합산을 빼먹어도 값이 안 변해 검사가 무의미해진다.
TREE = {"e-영업1": ["가방", "지갑"], "e-영업2": ["슈즈", "가방"]}


def synth_orgcat():
    """조직 > 카테고리 2단 — 브랜드·상품 칸은 지금 실제 export처럼 비워 둔다."""
    rows = []

    def add(gran, label, sortkey, year, org, cat, met, v):
        rows.append({"gran": gran, "metric": met, "org": org, "cat": cat,
                     "brand": "", "item": "", "lfms": "N", "year": year,
                     "label": label, "close": "final", "sortkey": sortkey,
                     "value": float(v)})
    def node(gran, label, sk, y, org, cat, rev, cust):
        """한 노드에 다섯 지표를 다 심는다.

        `거래액 = 상품UV × 상품CR × 객단가`가 **정확히** 성립하게 만든다 — 실파일에서
        오차 0.000%로 맞는 항등식이라, 여기서 깨 두면 요인 분해 테스트가 아무것도 못 잡는다.
        """
        uv, cr = cust * 20.0, 0.05
        add(gran, label, sk, y, org, cat, "첫구매 거래액", rev)
        add(gran, label, sk, y, org, cat, "첫구매 고객수", cust)
        add(gran, label, sk, y, org, cat, "첫구매 객단가", rev / cust)
        add(gran, label, sk, y, org, cat, "상품UV", uv)
        add(gran, label, sk, y, org, cat, "상품CR", cr)

    for y in YEARS:
        # 거래액과 고객수를 서로 다른 비율로 흔들어 객단가도 같이 움직이게 한다 —
        # 둘이 같은 배수면 객단가가 고정돼 요인 분해가 한 요인만 가리킨다
        kr, kc = (1.0, 1.0) if y == 2026 else (0.8, 0.9)
        for gran, labels in (("월", [(f"{m}월", m * 100) for m in MONTHS]),
                             ("주", [(f"{m:02d}월 {w}주차", m * 100 + w)
                                     for m in MONTHS for w in WEEKS])):
            for label, sk in labels:
                sv = y * 10000 + sk
                node(gran, label, sv, y, "*TOTAL", "*TOTAL", 50e6 * kr, 500 * kc)
                base = 30e6
                for org, cats in TREE.items():
                    node(gran, label, sv, y, org, "*TOTAL", base * kr, 300 * kc)
                    part = base / len(cats)
                    for cat in cats:
                        node(gran, label, sv, y, org, cat, part * kr, 150 * kc)
                    base = 20e6
    return pd.DataFrame(rows)[W.ORGCAT_COLS]


def _open(store=None, orgcat=None, mode=None):
    """임시 폴더에 스토어를 깔고 앱을 띄운 뒤 새 페이지로 이동한다."""
    from streamlit.testing.v1 import AppTest
    tmp = tempfile.mkdtemp()
    shutil.copy(APP, os.path.join(tmp, "weekly_report.py"))
    _te = ROOT / "table_export.py"
    if _te.exists():
        shutil.copy(_te, os.path.join(tmp, "table_export.py"))
    (synth_store() if store is None else store).to_csv(
        os.path.join(tmp, W.DATA_STORE), index=False, encoding="utf-8-sig")
    oc = synth_orgcat() if orgcat is None else orgcat
    if oc is not None and not oc.empty:
        oc.to_csv(os.path.join(tmp, W.ORGCAT_STORE), index=False, encoding="utf-8-sig")
    # 스토어가 **cwd 기준 CSV**라 여기서 나오면 안 된다. 테스트가 끝날 때까지 임시 폴더에
    # 머물고, 되돌리는 건 러너(main)가 맡는다 — 중간에 돌아오면 다음 at.run()부터
    # 스토어가 통째로 사라져 '누적 데이터가 없어요' 화면을 보고 통과해 버린다.
    os.chdir(tmp)
    at = AppTest.from_file(os.path.join(tmp, "weekly_report.py"), default_timeout=TIMEOUT)
    at.run()
    assert not at.exception, at.exception[0].value
    rad = [r for r in at.sidebar.radio if r.label == "페이지"]
    assert rad, "페이지 라디오를 못 찾았어요"
    rad[0].set_value(PAGE); at.run()
    assert not at.exception, at.exception[0].value
    if mode:
        cmp_r = [r for r in at.radio if r.label == "비교 기준"]
        assert cmp_r, f"비교 기준 라디오가 없어요 — {[r.label for r in at.radio]}"
        cmp_r[0].set_value(mode); at.run()
        assert not at.exception, at.exception[0].value
    return at


def _html(at):
    return "\n".join(str(m.value) for m in at.markdown)


def _texts(at):
    out = []
    for coll in (at.markdown, at.caption, at.info, at.warning, at.success,
                 at.title, at.header, at.subheader):
        out += [str(e.value) for e in coll]
    return out


def _kpi(at, label):
    m = re.search(r'kpi-label">' + re.escape(label) + r'</div><div class="kpi-value">([^<]*)</div>',
                  _html(at))
    return m.group(1) if m else None


def _rate(at, label):
    m = re.search(r'>' + re.escape(label) + r'</div>\s*<div[^>]*font-weight:700[^>]*>([^<]*)</div>',
                  _html(at))
    return m.group(1) if m else None


def _tail(at, label):
    m = re.search(r'>' + re.escape(label) + r' <b[^>]*>([^<]*)</b>', _html(at))
    return m.group(1) if m else None


def _frames(at):
    out = []
    for t in at.dataframe:
        v = getattr(t.value, "data", t.value)
        if hasattr(v, "columns"):
            out.append(v)
    return out


def _org_frames(at):
    """④의 조직 표. 이름이 **일반 칼럼**이라 index.name으로는 못 찾는다."""
    return [f for f in _frames(at) if "조직" in f.columns]


def _clickable(at):
    """행 클릭으로 파고드는 표들 — 셀 선택 + 인덱스 숨김이어야 한다."""
    out = []
    for t in at.dataframe:
        kw = getattr(t, "proto", None)
        key = getattr(t, "key", None) or ""
        if "wr_fn_orgsel" in str(key) or "wr_fn_catsel" in str(key):
            out.append(t)
    return out


def _sel_step(at):
    """③ 채널별 증감의 지표 셀렉트 — 퍼널 단계가 선택지에 있는 쪽."""
    got = [x for x in at.selectbox if x.label == "지표" and "가입자수" in x.options]
    assert got, [(x.label, list(x.options)) for x in at.selectbox]
    return got[0]


def _sel_oc(at):
    """④ 조직·카테고리의 지표 셀렉트 — 첫구매 지표만 있는 쪽."""
    got = [x for x in at.selectbox if x.label == "지표" and "가입자수" not in x.options]
    assert got, [(x.label, list(x.options)) for x in at.selectbox]
    return got[0]


CASES = []


def case(fn):
    CASES.append(fn)
    return fn


# ── 등록·렌더 ───────────────────────────────────────────────────────
@case
def t_page_is_registered_second():
    """요약 바로 다음 자리여야 한다 — 개괄에서 이 화면으로 내려오는 흐름이라서."""
    at = _open()
    opts = list([r for r in at.sidebar.radio if r.label == "페이지"][0].options)
    assert PAGE in opts, f"페이지가 없어요 — {opts}"
    assert opts.index(PAGE) == 1, f"두 번째 자리가 아니에요 — {opts}"


@case
def t_all_blocks_render():
    at = _open()
    txt = " ".join(_texts(at))
    for mark in ("① ", "② ", "③ ", "④ ", "⑤ "):
        assert mark in txt, f"{mark} 블록이 안 보여요"


@case
def t_mtd_mode_renders():
    at = _open(mode="월누적(MTD) — 전년 동월")
    txt = " ".join(_texts(at))
    assert "① " in txt and "④ " in txt, "MTD 모드에서 블록이 빠졌어요"
    assert _kpi(at, "가입자수"), "MTD 모드에서 퍼널 카드가 안 그려졌어요"


@case
def t_funnel_left_page_01():
    """01에 그대로 남아 있으면 같은 화면이 두 벌이 된다."""
    at = _open()
    [r for r in at.sidebar.radio if r.label == "페이지"][0].set_value("01. 주간보고 요약")
    at.run()
    assert not at.exception, at.exception[0].value
    txt = " ".join(_texts(at))
    assert "전환 퍼널" not in txt, "01에 전환 퍼널이 남아 있어요"
    assert "채널 기여 분해" not in txt, "01에 채널 기여 분해가 남아 있어요"


# ── 비율 두 칸의 출처 ───────────────────────────────────────────────
@case
def t_join_rate_is_computed_not_read_from_file():
    """가입율은 `가입자수 ÷ 비회원트래픽`이다.

    파일의 「가입율」 지표를 50%로 심어 뒀다 — 그 값이 화면에 뜨면 계산이 아니라 파일을
    읽고 있다는 뜻이고, 채널로 쪼갤 때 합계와 규칙이 갈린다.
    """
    at = _open()
    got = _rate(at, "가입율")
    assert got == "1.00%", f"가입율이 계산값이 아니에요 — {got} (파일 값이면 50.00%)"


@case
def t_daily_cr_comes_from_the_file():
    """당일가입CR은 분자가 데이터에 없어 역산이 불가능하다 — 파일 값 그대로여야 한다."""
    at = _open()
    got = _rate(at, "당일가입 첫구매율")
    assert got == f"{DAILY_CR * 100:.2f}%", f"당일가입CR이 파일 값이 아니에요 — {got}"


@case
def t_aov_keeps_the_file_value():
    """객단가는 「01」 KPI 카드와 같은 숫자여야 한다 — 거래액÷고객수로 다시 만들지 않는다."""
    at = _open()
    got = _tail(at, "객단가")
    assert got == f"{FILE_AOV:,}원", f"객단가가 파일 값이 아니에요 — {got} (역산이면 100,000원)"


@case
def t_funnel_does_not_pretend_to_multiply_out():
    """`가입자수 × 당일가입CR ≠ 첫구매 고객수`라는 걸 화면이 밝혀야 한다.

    합성값으로 1,000 × 7.25% = 72.5명인데 첫구매 고객수는 500명이다. 안 밝히면
    보는 사람이 산식이 맞는 줄 알고 읽는다.
    """
    at = _open()
    assert _kpi(at, "첫구매 고객수") == "500명", _kpi(at, "첫구매 고객수")
    txt = " ".join(_texts(at))
    assert "전체 첫구매" in txt and "안 맞아요" in txt, "곱셈이 안 닫힌다는 안내가 없어요"


@case
def t_counts_and_yoy_are_right():
    at = _open()
    assert _kpi(at, "비회원트래픽") == "100,000명", _kpi(at, "비회원트래픽")
    assert _kpi(at, "가입자수") == "1,000명", _kpi(at, "가입자수")
    # 가입자수 500 → 1,000 = +100.0%
    assert "+100.0% (전년동주)" in _html(at), "가입자수 전년비가 틀려요"


# ── ② 채널별 상세 ───────────────────────────────────────────────────
@case
def t_channel_table_covers_every_step():
    at = _open()
    fr = [f for f in _frames(at) if "비회원트래픽" in f.columns]
    assert fr, f"채널 표가 없어요 — {[list(f.columns) for f in _frames(at)]}"
    tbl = fr[0]
    for stp in W.FUNNEL_STEPS:
        assert stp in tbl.columns, f"{stp} 칸이 없어요"
        assert f"{stp} 전년비" in tbl.columns, f"{stp} 전년비 칸이 없어요"
    assert list(tbl.index)[:3] == ["전체", "직접", "광고"], list(tbl.index)


@case
def t_channel_rates_are_per_channel():
    """채널별 가입율은 **그 채널의 두 카운트**로 만들어야 한다.

    전체 값을 그대로 복사하면 세 줄이 전부 1.00%가 되어 '어느 채널이 새는지'를 못 본다.
    """
    at = _open()
    tbl = [f for f in _frames(at) if "비회원트래픽" in f.columns][0]
    got = {ix: tbl.loc[ix, "가입율"] for ix in ("전체", "직접", "광고")}
    assert got == {"전체": "1.00%", "직접": "2.00%", "광고": "0.33%"}, got


@case
def t_channel_table_carries_the_rate_caveat():
    at = _open()
    assert any("채널 합이 전체와 다른 게 정상" in t for t in _texts(at)), \
        "비율 지표 주의 문구가 없어요"


# ── ③ 채널 기여 분해 ────────────────────────────────────────────────
@case
def t_decomposition_offers_additive_only():
    """비율 지표는 채널 합 ≠ 전체라 분해가 성립하지 않는다 — 선택지에 있으면 안 된다."""
    at = _open()
    opts = list(_sel_step(at).options)
    for good in ("비회원트래픽", "가입자수", "첫구매 고객수", "첫구매 거래액"):
        assert good in opts, f"«{good}»가 빠졌어요 — {opts}"
    # 가산 지표를 고르면 워터폴(부분의 합 = 전체)이 그려진다
    _sel_step(at).set_value("첫구매 거래액"); at.run()
    assert not at.exception, at.exception[0].value
    assert any("증감 분해" in t for t in _texts(at)), "가산 지표인데 분해 설명이 없어요"


@case
def t_decomposition_draws_a_waterfall():
    at = _open()
    assert at.get("plotly_chart"), "워터폴 차트가 없어요"


# ── ④ 조직 > 카테고리 ───────────────────────────────────────────────
@case
def t_orgcat_lists_orgs_with_rowclick_hint():
    """조직 표는 늘 보이고, 한 단계 더는 **행을 눌러** 들어간다."""
    at = _open()
    assert not [s for s in at.selectbox if s.label == "1. 조직"], \
        "셀렉트박스가 남아 있어요 — 행 클릭으로 바뀌었어요"
    fr = _org_frames(at)
    assert fr, f"조직 표가 없어요 — {[list(f.columns) for f in _frames(at)]}"
    # 이름은 인덱스가 아니라 「조직」 칼럼에 있다(셀 선택이 인덱스 칸을 안 돌려준다)
    assert set(TREE) <= set(fr[0]["조직"]), list(fr[0]["조직"])
    assert any("누르면" in t and "카테고리" in t for t in _texts(at)), "행 클릭 안내가 없어요"


@case
def t_org_row_opens_its_categories():
    """행을 누른 뒤 나오는 표 — 화면이 부르는 것과 같은 헬퍼로 값을 대조한다."""
    oc = synth_orgcat()
    sub = oc[(oc["gran"] == "월") & (oc["lfms"] == "N")]
    view = W.orgcat_view(sub)
    kids = view.live(("e-영업1",))[0]
    assert set(kids) == set(TREE["e-영업1"]), kids
    tbl = W._funnel_level_table(view, ["e-영업1"], kids, "카테고리", "첫구매 거래액",
                                2026, 2025, "1월", "final",
                                view.get(("e-영업1",), "첫구매 거래액", 2026, "1월", "mtd"))
    # 이름은 인덱스가 아니라 칼럼이다 — 셀 선택이 인덱스 칸을 안 돌려주기 때문
    assert list(tbl["카테고리"]) == kids, list(tbl["카테고리"])
    assert "비중" in tbl.columns, list(tbl.columns)      # 거래액은 가산 지표
    # 가방·지갑이 각각 절반 → 비중 50.0%
    assert set(tbl["비중"]) == {"50.0%"}, tbl["비중"].tolist()


@case
def t_orgcat_share_only_for_additive():
    """하위 합이 상위와 맞는 건 거래액뿐이다 — 고객수에 비중을 붙이면 거짓말이 된다."""
    at = _open()
    _sel_oc(at).set_value("첫구매 거래액"); at.run()
    fr = _org_frames(at)[0]
    assert "비중" in fr.columns, f"거래액엔 비중이 있어야 해요 — {list(fr.columns)}"
    _sel_oc(at).set_value("첫구매 고객수"); at.run()
    fr = _org_frames(at)[0]
    assert "비중" not in fr.columns, f"고객수엔 비중이 없어야 해요 — {list(fr.columns)}"
    assert any("하위 합이 상위와 안 맞는" in t for t in _texts(at)), "왜 뺐는지 안 밝혔어요"


@case
def t_orgcat_missing_source_explains_itself():
    at = _open(orgcat=pd.DataFrame())
    assert any("조직×카테고리 데이터가 없어요" in t for t in _texts(at)), \
        "원천이 없을 때 안내가 없어요"
    assert not at.exception, at.exception[0].value


# ── ⑤ 앱 수신동의 ───────────────────────────────────────────────────
@case
def t_push_period_avg_reads_daily_rows():
    """앱푸시는 일별로만 쌓인다 — 월은 정확히, 주는 근사로 묶는다."""
    d = synth_store()
    v, approx, n = W.push_period_avg(d, 2026, "9월")
    assert abs(v - np.mean(list(PUSH_DAILY.values()))) < 1e-6, v
    assert approx is False, "월은 근사가 아니에요"
    assert n == len(PUSH_DAILY), f"며칠치인지 같이 돌려줘야 해요 — {n}"
    v2, approx2, n2 = W.push_period_avg(d, 2026, "09월 1주차")
    assert approx2 is True, "주는 근사라고 알려야 해요"
    assert abs(v2 - PUSH_DAILY[1]) < 1e-6, v2      # 1~7일 = 1일치뿐
    assert n2 == 1, n2
    assert pd.isna(W.push_period_avg(d, 2026, "없는라벨")[0])


@case
def t_app_block_shows_consent_rate():
    at = _open()
    fr = [f for f in _frames(at) if f.index.name == "비율"]
    assert fr, f"앱 비율 표가 없어요 — {[f.index.name for f in _frames(at)]}"
    tbl = fr[0]
    # 앱 보유 무관 원천이 없으면 이름은 예전 그대로다(구분할 대상이 하나뿐이라서)
    assert "신규회원 앱 수신동의율" in tbl.index, list(tbl.index)
    # 주간 기본 모드: 그 주 앱푸시 일평균 ÷ 가입자수
    assert tbl.loc["신규회원 앱 수신동의율", "2026년"] != "–", tbl.to_dict()


@case
def t_missing_app_install_says_why():
    """앱설치 원천이 없으면 빈 줄만 두지 말고 왜 비었는지 말해야 한다."""
    at = _open()
    assert any("앱설치" in t and "안 올라와서" in t for t in _texts(at)), \
        "앱설치 부재 안내가 없어요"


@case
def t_app_install_file_is_recognized():
    """올릴 예정인 원천 — 파일명만으로 지표가 잡혀야 한다."""
    for nm in ("월_앱설치(일평균).xlsx", "주_신규앱설치.xlsx", "일_앱다운로드.csv"):
        kind, gran, met = W.detect_file(nm)
        assert (kind, met) == ("metric", "앱설치"), f"{nm} → {(kind, gran, met)}"
    assert W.METRIC_UNIT["앱설치"] == ("명", 1)


# ══════════════════════════════════════════════════════
# 앱설치 원천 (LFmall 앱 대시보드 export)
# ══════════════════════════════════════════════════════
# 실파일에서 그대로 뽑은 조각이다. 같은 대시보드가 세 벌을 따로 떨어뜨리는데 **모양이
# 다 다르다** — 일별은 연도가 없고, 주간은 주가 일~토고, 월간은 칼럼이 통째로 다르다.
_AI_HDR = "날짜,전체 설치,신규 설치,재설치,스토어 방문,삭제,Push 활성 기기"
# (표시날짜, 전체, 신규, 재설치, 스토어, 삭제, Push활성) — 실파일 값 그대로
_AI_DAILY = [
    ("9월 5일", 710, 394, 316, 1080, "", 547529),
    ("9월 4일", 615, 339, 276, 1027, "", 547371),
    ("9월 3일", 721, 397, 324, 1241, "", 547413),
    ("9월 2일", 815, 439, 376, 1311, "", 547491),
    ("9월 1일", 3269, 2147, 1122, 4410, "", 1124697),
    ("8월 31일", 819, 447, 372, 1442, "", 546997),
    ("8월 30일", 2456, 1549, 907, 3105, 1685, 1122732),
    ("8월 29일", 3024, 2171, 853, 3005, 1532, 1122054),
    ("8월 28일", 3256, 2369, 887, 2724, 1669, 1121375),
    ("8월 27일", 5073, 3705, 1368, 4859, 1927, 1120345),
    ("8월 26일", 2811, 2076, 735, 2315, 1346, 1121447),
    ("8월 25일", 2846, 2000, 846, 2858, 1500, 1121494),
    ("8월 24일", 2916, 2036, 880, 2864, 1677, 1121088),
    ("8월 23일", 2862, 1910, 952, 2854, 1697, 1120921),
    ("3월 31일", 2171, 1231, 940, 2615, 1303, 1123067),
    ("3월 30일", 1952, 1234, 718, 2521, 1136, 1123613),
    ("3월 1일", 2229, 1232, 997, 2347, 1386, 1113894),
]
# 주간 파일 — 주가 **일~토**다. `~ 26.08.29`는 8/23~8/29(합 22,788)이고
# 보고서의 「08월 4주차」는 8/24~8/30(합 22,382)이라 **다른 창**이다.
_AI_WEEKLY = "\n".join([_AI_HDR,
                        "~ 26.09.05,5337,2895,2442,8719,,547274",
                        "~ 26.08.29,22788,16267,6521,21479,11348,1121246",
                        "~ 26.08.22,20169,13729,6440,19524,11044,1119307"])
# 월간 파일 — 칼럼이 다르고, 실파일에서 최신 달이 Android 0으로 깨져 온다
# (2026-08: 파일 19,509 vs 일별 합 92,696).
_AI_MONTHLY = "\n".join(["날짜,전체 설치 (전체),전체 설치 (Android),전체 설치 (iOS)",
                         "2026. 08,19509,0,19509",
                         "2026. 07,84071,65137,18934",
                         "2026. 06,69744,57794,11950"])
_AI_WEEK_SUN_SAT = 22788        # 주간 파일이 준 8/23~8/29
_AI_WEEK_MON_SUN = 22382        # 보고서 규칙의 08월 4주차 (8/24~8/30)


def appinstall_csv(rows=None, hdr=_AI_HDR):
    out = [hdr]
    for r in (rows or _AI_DAILY):
        out.append(",".join("" if x == "" else str(x) for x in r))
    return "\n".join(out).encode("utf-8")


def synth_appinstall_store(per_day=3000, days=250, broken_tail=0):
    """기간을 안 가리게 **모든 날 같은 값**으로 깐다 — 어느 주·달을 골라도 일평균이 같다.

    `broken_tail`이 있으면 마지막 N일의 기기 수를 반토막 낸다 (실파일의 한쪽 플랫폼
    누락과 같은 모양).
    """
    rows, d = [], datetime.date(2026, 9, 5)
    for i in range(days):
        bad = i < broken_tail
        rows.append((f"{d.month}월 {d.day}일",
                     per_day // 4 if bad else per_day,
                     (per_day - 800) // 4 if bad else per_day - 800, 800,
                     per_day + 200, 1000, 500_000 if bad else 1_100_000))
        d -= datetime.timedelta(days=1)
    return W.parse_appinstall_file("일별.csv", appinstall_csv(rows))


@case
def t_appinstall_daily_is_recognized():
    d = W.parse_appinstall_file("일별.csv", appinstall_csv())
    assert d.attrs.get("appinstall_kind") == "일", d.attrs
    assert set(d["gran"]) == {"일", "주", "월"}, set(d["gran"])
    assert "앱설치" in set(d["metric"]), sorted(set(d["metric"]))
    assert set(d.columns) == set(W.STORE_COLS), list(d.columns)
    assert d.attrs["date_range"] == ("2026-03-01", "2026-09-05"), d.attrs["date_range"]


@case
def t_daily_year_is_inferred_backwards():
    """일별 파일엔 **연도가 없다.** 최근 행을 오늘 기준으로 잡고 거슬러 올라간다."""
    got = W._ai_days(["1월 3일", "1월 1일", "12월 31일", "12월 28일"],
                     today=datetime.date(2026, 1, 5))
    assert got == [datetime.date(2026, 1, 3), datetime.date(2026, 1, 1),
                   datetime.date(2025, 12, 31), datetime.date(2025, 12, 28)], got
    # 오름차순으로 와도 같은 답이 나와야 한다
    assert W._ai_days(["12월 28일", "12월 31일", "1월 1일", "1월 3일"],
                      today=datetime.date(2026, 1, 5)) == got[::-1]
    # 아직 안 온 날짜면 작년으로 — 12월 파일을 1월에 받는 경우
    assert W._ai_days(["12월 20일"], today=datetime.date(2026, 1, 5)) == \
        [datetime.date(2025, 12, 20)]


@case
def t_week_labels_follow_the_report_rule():
    """주차는 목요일이 속한 달 기준이라 달을 넘나든다 — 마스터와 같은 규칙이어야 조인된다."""
    d = W.parse_appinstall_file("일별.csv", appinstall_csv())
    wk = set(d[d["gran"] == "주"]["label"])
    assert "09월 1주차" in wk, wk          # 8/31(월) → 목요일 9/3 → 9월 1주차
    assert "04월 1주차" in wk, wk          # 3/30(월) → 목요일 4/2 → 4월 1주차
    mo = set(d[d["gran"] == "월"]["label"])
    assert "3월" in mo and "4월" not in mo, mo    # 달은 달력 기준 (3/30·31은 3월)


@case
def t_values_are_daily_means_not_sums():
    """다른 값이 전부 일평균이라 합계로 내면 비율이 통째로 틀어진다."""
    d = W.parse_appinstall_file("일별.csv", appinstall_csv())
    week = [r for r in _AI_DAILY if r[0].startswith("8월 2") and int(r[0][3:-1]) >= 24]
    want = np.mean([r[1] for r in week] + [2456])          # 8/24~8/30 전체 설치
    got = d[(d["gran"] == "주") & (d["label"] == "08월 4주차")
            & (d["metric"] == "앱_전체설치")]["value"]
    assert len(got) == 1 and abs(float(got.iloc[0]) - want) < 1e-6, (got.tolist(), want)
    assert abs(float(got.iloc[0]) * 7 - _AI_WEEK_MON_SUN) < 1e-6, float(got.iloc[0]) * 7
    # `앱설치`는 **신규 설치** 칸이다 — 전체 설치를 집으면 재설치가 섞여 값이 커진다
    nw = np.mean([r[2] for r in week] + [1549])
    gn = d[(d["gran"] == "주") & (d["label"] == "08월 4주차")
           & (d["metric"] == "앱설치")]["value"]
    assert len(gn) == 1 and abs(float(gn.iloc[0]) - nw) < 1e-6, (gn.tolist(), nw)
    assert float(gn.iloc[0]) < float(got.iloc[0]), "신규 < 전체 여야 해요"


@case
def t_blank_deletes_stay_empty_not_zero():
    """8/31부터 `삭제` 칸이 통째로 비어 온다 — 0으로 채우면 삭제가 없었던 게 된다."""
    d = W.parse_appinstall_file("일별.csv", appinstall_csv())
    wk = d[(d["gran"] == "주") & (d["label"] == "09월 1주차")]
    assert "앱설치" in set(wk["metric"]), set(wk["metric"])
    assert "앱_삭제" not in set(wk["metric"]), "빈 삭제가 값으로 들어왔어요"
    aug = d[(d["gran"] == "주") & (d["label"] == "08월 4주차") & (d["metric"] == "앱_삭제")]
    assert len(aug) == 1, "값이 있는 주는 그대로 있어야 해요"


@case
def t_weekly_file_is_recognized_but_not_stored():
    """주간 export는 주가 **일~토**라 보고서의 월~일과 다른 창이다.

    그런데 목요일로 라벨을 매기면 **같은 라벨**이 나와서, 저장하면 일별에서 만든 값을
    조용히 덮어쓴다. 두 값이 실제로 다르다는 걸 같이 박아 둔다.
    """
    assert _AI_WEEK_SUN_SAT != _AI_WEEK_MON_SUN, "전제가 깨졌어요"
    d = W.parse_appinstall_file("주간.csv", _AI_WEEKLY.encode("utf-8"))
    assert d.empty, "주간 파일이 저장됐어요 — 일별 값을 덮어써요"
    assert d.attrs.get("appinstall_kind") == "주", d.attrs
    cls = W.classify_uploads((("주간.csv", _AI_WEEKLY.encode("utf-8")),))
    assert any("앱설치(주간)" in c[1] and "저장 안 함" in c[1] for c in cls), cls


@case
def t_monthly_file_is_recognized_but_not_stored():
    """월간 export는 칼럼이 다르고 최신 달이 깨져 온다(Android 0)."""
    d = W.parse_appinstall_file("월간.csv", _AI_MONTHLY.encode("utf-8"))
    assert d.empty and d.attrs.get("appinstall_kind") == "월", d.attrs
    cls = W.classify_uploads((("월간.csv", _AI_MONTHLY.encode("utf-8")),))
    assert any("앱설치(월간)" in c[1] and "저장 안 함" in c[1] for c in cls), cls


@case
def t_master_parser_does_not_swallow_them():
    """세 파일 다 마스터·앱푸시 파서로 새면 안 된다."""
    pf, d = W.route_push("일별.csv", appinstall_csv())
    assert pf is None and d is not None and not d.empty, (pf, d)
    assert set(d["metric"]) <= set(W.APPINSTALL_METS) | {W.APPINSTALL_FLAG}, \
        sorted(set(d["metric"]))
    for nm, raw in (("주간.csv", _AI_WEEKLY), ("월간.csv", _AI_MONTHLY)):
        pf, d = W.route_push(nm, raw.encode("utf-8"))
        assert pf is None and d is None, f"{nm} → {(pf, d)}"


@case
def t_unreadable_dates_are_counted_not_swallowed():
    """날짜를 하나도 못 읽으면 '왜 비었는지'까지 말해야 한다."""
    bad = appinstall_csv([("알수없음",) + r[1:] for r in _AI_DAILY])
    d = W.parse_appinstall_file("깨진일별.csv", bad)
    assert d.empty, "못 읽는 날짜인데 값이 나왔어요"
    cls = W.classify_uploads((("깨진일별.csv", bad),))
    assert any("앱설치" in c[1] and "못 읽었" in c[1] for c in cls), cls


@case
def t_recognized_list_names_the_range():
    """읽은 범위를 찍어 둬야 연도를 잘못 잡은 걸 눈으로 잡는다."""
    cls = W.classify_uploads((("일별.csv", appinstall_csv()),))
    assert any("✅ 앱설치 원천(일별)" in c[1] and "2026-03-01~2026-09-05" in c[1]
               for c in cls), cls


@case
def t_appinstall_reaches_the_funnel_page():
    store = pd.concat([synth_store(), synth_appinstall_store(3000)], ignore_index=True)
    # 앱설치는 **실제 달력 주**라 마스터 합성본의 `09월 4주차`(9/21~)와 안 겹친다 —
    # 월 비교로 본다(9월 = 9/1~9/5, 값이 모든 날 같아 일평균은 그대로 3,000).
    at = _open(store=store, mode="월누적(MTD) — 전년 동월")
    # synth: 전체 3,000 = 신규 2,200 + 재설치 800 → 카드는 **신규**를 보여야 한다
    assert _kpi(at, "앱설치") == "2,200명", _kpi(at, "앱설치")
    fr = [f for f in _frames(at) if f.index.name == "비율"][0]
    assert fr.loc["가입자 대비 앱 신규설치율", "2026년"] != "–", fr.to_dict()
    assert any("앱설치 상세" in str(e.label) for e in at.expander), \
        [str(e.label) for e in at.expander]


@case
def t_prior_year_absence_is_explained():
    """전년이 없으면 빈 칸만 두지 말고 왜 비었는지 말해야 한다 (실제로 전년이 없다)."""
    store = pd.concat([synth_store(), synth_appinstall_store(3000)], ignore_index=True)
    at = _open(store=store, mode="월누적(MTD) — 전년 동월")
    assert any("전년 데이터가 없어" in t for t in _texts(at)), "전년 부재 안내가 없어요"

@case
def t_broken_collection_days_are_flagged():
    """`Push 활성 기기`는 그날의 **잔고**라 하루 만에 반토막 났다가 돌아올 수 없다.

    실파일에서 8/31·9/2~9/5가 그렇게 온다(한쪽 플랫폼 누락). 잔고가 정상인 9/1은
    같은 구간에 있어도 잡히면 안 된다 — 잡히면 멀쩡한 날까지 의심하게 된다.
    """
    d = W.parse_appinstall_file("일별.csv", appinstall_csv())
    bad = d.attrs.get("level_drop") or []
    assert bad == ["2026-08-31", "2026-09-02", "2026-09-03",
                   "2026-09-04", "2026-09-05"], bad
    day = d[(d["gran"] == "일") & (d["metric"] == W.APPINSTALL_FLAG)]
    assert set(day["label"]) == {"8/31", "9/2", "9/3", "9/4", "9/5"}, set(day["label"])
    assert "9/1" not in set(day["label"]), "잔고가 정상인 9/1이 잡혔어요"
    wk = d[(d["gran"] == "주") & (d["metric"] == W.APPINSTALL_FLAG)]
    assert dict(zip(wk["label"], wk["value"])) == {"09월 1주차": 5.0}, wk.to_dict("records")
    # 이상 없는 기간엔 아예 안 생긴다 (0으로 깔면 '표시'가 아니라 소음이 된다)
    assert "03월 2주차" not in set(wk["label"]), set(wk["label"])


@case
def t_broken_days_are_not_corrected():
    """표시만 한다 — 값을 지우거나 고치면 원천과 화면이 갈린다."""
    d = W.parse_appinstall_file("일별.csv", appinstall_csv())
    raw = {r[0]: r[2] for r in _AI_DAILY}          # 신규 설치 칸
    for lab, key in (("9/2", "9월 2일"), ("8/31", "8월 31일")):
        v = d[(d["gran"] == "일") & (d["label"] == lab) & (d["metric"] == "앱설치")]["value"]
        assert len(v) == 1 and float(v.iloc[0]) == raw[key], (lab, v.tolist(), raw[key])


@case
def t_broken_days_warn_on_screen():
    """숫자를 보는 자리에서 말해야 한다 — 인식 목록은 업로드 때 한 번 스쳐 간다."""
    store = pd.concat([synth_store(), synth_appinstall_store(3000, broken_tail=3)],
                      ignore_index=True)
    at = _open(store=store, mode="월누적(MTD) — 전년 동월")
    warn = [str(w.value) for w in at.warning]
    assert any("절반 이하로 찍힌 날이 3일" in w for w in warn), warn
    assert any("값은 원천 그대로 두었어요" in w for w in warn), warn


@case
def t_clean_data_does_not_warn():
    """멀쩡한 데이터에 경고가 뜨면 정작 진짜 문제를 무시하게 된다."""
    store = pd.concat([synth_store(), synth_appinstall_store(3000)], ignore_index=True)
    at = _open(store=store, mode="월누적(MTD) — 전년 동월")
    assert not any("절반 이하" in str(w.value) for w in at.warning), \
        [str(w.value) for w in at.warning]


@case
def t_recognized_list_mentions_broken_days():
    cls = W.classify_uploads((("일별.csv", appinstall_csv()),))
    assert any("기기 수가 절반 이하인 날 5일" in c[1] and "08-31" in c[1] for c in cls), cls


@case
def t_category_rollup_spans_orgs():
    """조직을 가로질러 카테고리를 모은 표 — 조직을 하나씩 안 들어가도 보인다."""
    at = _open()
    fr = [f for f in _frames(at) if f.index.name == "카테고리"]
    assert fr, f"카테고리 합산 표가 없어요 — {[list(f.columns) for f in _frames(at)]}"
    cats = set(fr[0].index)
    want = {c for v in TREE.values() for c in v}
    assert cats == want, (cats, want)


@case
def t_category_rollup_sums_only_what_can_be_summed():
    """거래액은 조직 합, 객단가는 거래액합÷고객수합. 고객수 합은 중복이 섞인다고 밝힌다."""
    at = _open()
    _sel_oc(at).set_value("첫구매 거래액"); at.run()
    fr = [f for f in _frames(at) if f.index.name == "카테고리"][0]
    assert "비중" in fr.columns, list(fr.columns)
    assert any("조직 합이 전체와 맞는" in t for t in _texts(at)), "거래액 설명이 없어요"

    _sel_oc(at).set_value("첫구매 객단가"); at.run()
    assert not at.exception, at.exception[0].value
    fr = [f for f in _frames(at) if f.index.name == "카테고리"][0]
    assert "비중" not in fr.columns, "객단가는 더할 수 없어 비중을 붙이면 안 돼요"
    assert any("거래액 합 ÷ 고객수 합" in t for t in _texts(at)), "객단가 산식 설명이 없어요"


@case
def t_category_rollup_carries_uv_and_cr():
    """상품UV·상품CR도 조직 합산 표에 나온다.

    원천이 조직×카테고리로 주는 값이라 낱개는 그대로 읽고, 가로지를 땐 UV는 더하고
    CR은 **고객수 합 ÷ 상품UV 합**으로 되만든다(`거래액 = 상품UV × 상품CR × 객단가`와
    `거래액 = 고객수 × 객단가`에서 나오는 항등식). 더하기로 되돌리면 CR이 카테고리
    개수만큼 부풀어 오른다.
    """
    at = _open()
    for met in ("상품UV", "상품CR"):
        _sel_oc(at).set_value(met); at.run()
        assert not at.exception, at.exception[0].value
        fr = [f for f in _frames(at) if f.index.name == "카테고리"]
        assert fr, f"«{met}» 합산 표가 없어요 — {[list(f.columns) for f in _frames(at)]}"
        cats = set(fr[0].index)
        want = {c for v in TREE.values() for c in v}
        assert cats == want, (met, cats, want)
        # 둘 다 유니크/비율이라 비중 칸이 붙으면 안 된다
        assert "비중" not in fr[0].columns, f"«{met}»에 비중이 붙었어요 — {list(fr[0].columns)}"
    assert any("고객수 합 ÷ 상품UV 합" in t for t in _texts(at)), "상품CR 산식 설명이 없어요"


@case
def t_rollup_cr_matches_the_identity():
    """합산 상품CR은 **고객수 합 ÷ 상품UV 합**이다 — 화면 값을 직접 읽어 대조한다.

    손으로 계산한 값끼리만 맞춰 보면 화면이 그냥 더하고 있어도 통과한다. 그래서
    합산 표에서 값을 꺼내 온다. 「가방」은 두 조직이 나눠 가져서, 더하기로 되돌리면
    CR이 두 배로 부풀어 이 검사에 걸린다.
    """
    oc = synth_orgcat()
    sub = oc[(oc["gran"] == "월") & (oc["lfms"] == "N")]
    view = W.orgcat_view(sub)
    cat = next(c for c in {c for v in TREE.values() for c in v}
               if sum(c in v for v in TREE.values()) > 1)
    cu = uv = 0.0
    for o in view.live(())[0]:
        if cat not in view.live((o,))[0]:
            continue
        cu += view.get((o, cat), "첫구매 고객수", 2026, "1월", "mtd")
        uv += view.get((o, cat), "상품UV", 2026, "1월", "mtd")
    assert cu > 0 and uv > 0, (cat, cu, uv)
    want = W.fmt_value("상품CR", cu / uv)

    at = _open(mode="월누적(MTD) — 전년 동월")
    _sel_oc(at).set_value("상품CR"); at.run()
    assert not at.exception, at.exception[0].value
    fr = [f for f in _frames(at) if f.index.name == "카테고리"]
    assert fr, f"합산 표가 없어요 — {[list(f.columns) for f in _frames(at)]}"
    got = fr[0].loc[cat, "2026년"]
    assert got == want, f"«{cat}» 합산 상품CR이 {got} — 고객수합÷UV합이면 {want}"


@case
def t_org_table_is_click_to_drill_not_checkbox():
    """조직·카테고리 표는 **셀 선택**이어야 한다.

    `single-row`는 화면에 체크박스 열로 나와 행을 눌러도 안 잡힌다. 셀 선택이면
    아무 칸이나 눌러서 내려갈 수 있다. 다만 셀 선택은 **인덱스 칸을 안 돌려주므로**
    이름이 일반 칼럼이어야 하고 인덱스는 숨겨야 한다 — 셋이 한 묶음이다.
    """
    at = _open()
    hit = [t for t in at.dataframe if "wr_fn_orgsel" in str(getattr(t, "key", "") or "")]
    assert hit, "조직 표에 key가 없어요 — 행 클릭을 걸 데가 없어요"
    from streamlit.proto.Dataframe_pb2 import Dataframe as _DfP
    sp = hit[0].proto
    assert list(sp.selection_mode) == [_DfP.SelectionMode.SINGLE_CELL], \
        f"셀 선택이 아니에요(행 선택은 체크박스 열로 나와요) — {list(sp.selection_mode)}"
    # hide_index는 프로토 필드가 아니라 컬럼 설정 JSON으로 실린다
    assert json.loads(sp.columns or "{}").get("_index", {}).get("hidden") is True, \
        f"인덱스를 숨기지 않으면 셀 선택에 빈 칸이 하나 생겨요 — {sp.columns}"
    assert "조직" in _org_frames(at)[0].columns, "이름이 인덱스에 있으면 눌러도 안 잡혀요"


@case
def t_picked_row_gets_a_whole_row_background():
    """고른 행은 **줄 전체**에 배경이 들어와야 한다 — 셀 테두리만으론 안 보인다."""
    import pandas as _pd
    df = _pd.DataFrame({"조직": ["가", "나", "다"], "2026년": [1, 2, 3]})
    sty = W._hl_row(df.style, 1)
    ctx = sty._compute().ctx
    # ctx는 (행, 열) → 스타일 목록. 1번 행만 전 칼럼이 칠해져야 한다.
    painted = {r for (r, _c), v in ctx.items()
               if any("background-color" in str(x) for x in v)}
    assert painted == {1}, f"1번 행만 칠해져야 해요 — {painted}"
    cols = {c for (r, c), v in ctx.items()
            if r == 1 and any("background-color" in str(x) for x in v)}
    assert cols == {0, 1}, f"줄 전체가 아니라 일부만 칠해졌어요 — {cols}"
    # None이면 손대지 않는다
    assert W._hl_row(df.style, None)._compute().ctx == {}


@case
def t_highlight_does_not_leak_into_excel():
    """하이라이트는 화면 전용이다 — 받은 파일에 한 줄만 색칠돼 있으면 뭔지 알 수 없다.

    `wtable(dl_data=)`로 색 안 입힌 쪽을 내려받기에 넘기는지 본다.
    """
    src = inspect.getsource(W.funnel_orgcat_block if hasattr(W, "funnel_orgcat_block")
                            else W._funnel_level_table)
    whole = pathlib.Path(APP).read_text(encoding="utf-8")
    for key in ("wr_fn_orgsel", "_ckey"):
        i = whole.index(f"key={key}" if key.startswith("_") else f'key="{key}"')
        seg = whole[i - 400:i + 400]
        assert "dl_data=" in seg, f"«{key}» 표가 하이라이트를 엑셀까지 보내요"
    # wtable이 dl_data를 실제로 쓰는지
    assert "_src = data if dl_data is None else dl_data" in whole, \
        "wtable이 dl_data를 안 써요 — 넘겨도 소용없어요"


@case
def t_lower_blocks_start_collapsed():
    """하위 블록은 접혀 있다가 위를 누르면 펼쳐진다."""
    at = _open()
    labels = {str(e.label): e for e in at.expander}
    assert any("어디에서 빠졌나" in x for x in labels), f"요인 분해 접이식이 없어요 — {list(labels)}"
    assert any("카테고리별 (조직 합산)" in x for x in labels), \
        f"합산 접이식이 없어요 — {list(labels)}"
    # 아무것도 안 골랐으면 둘 다 접혀 있어야 한다
    for lb, e in labels.items():
        if "어디에서 빠졌나" in lb or "카테고리별 (조직 합산)" in lb:
            assert not e.proto.expanded, f"«{lb}»가 처음부터 펼쳐져 있어요"
    assert any("전체" in t and "누르면" in t for t in _texts(at)), \
        "왜 전체 기준인지 안 밝혔어요"


@case
def t_monthly_trend_puts_yoy_next_to_each_month():
    """월별 추이표 — 당해년도는 달마다 **바로 옆에** 전년 대비 증감이 붙는다.

    2025년 1월과 2026년 1월이 열두 칸 떨어져 있어 그냥은 못 맞댄다.
    """
    df = synth_store()
    tbl = W.trend_table(df, "월", W.METRICS7, [2025, 2026], delta_year=2026)
    cols = list(tbl.columns)
    cur = [c for c in cols if c[0] == 2026]
    # 값 칸 바로 다음이 그 달의 증감 칸이어야 한다
    plain = [c for c in cur if "증감" not in c[1]]
    for c in plain:
        i = cols.index(c)
        assert i + 1 < len(cols) and cols[i + 1] == (2026, f"{c[1]} 증감"), \
            f"«{c[1]}» 옆에 증감이 없어요 — {cols[i:i + 2]}"
    # 전년(2025)엔 안 붙는다 — 그 앞해가 없어 늘 '–'라 자리만 먹는다
    assert not [c for c in cols if c[0] == 2025 and "증감" in c[1]], \
        [c for c in cols if c[0] == 2025]
    # 값이 실제 전년비인지
    a = W.pick(df, "월", "첫구매 거래액", "*TOTAL", 2026, "1월", "final")
    b = W.pick(df, "월", "첫구매 거래액", "*TOTAL", 2025, "1월", "final")
    want = W.fmt_delta("첫구매 거래액", a, b)
    got = tbl.loc["첫구매 거래액", (2026, "1월 증감")]
    assert got == want, f"{got} vs {want}"


@case
def t_trend_delta_keeps_its_own_formatting():
    """증감 칸은 이미 문자열이다 — fmt_value를 다시 먹이면 뭉개진다."""
    df = synth_store()
    tbl = W.trend_table(df, "월", W.METRICS7, [2025, 2026], delta_year=2026)
    sty = W.style_trend(tbl, W.METRICS7)
    disp = getattr(sty, "data", sty)
    v = str(disp.loc["첫구매 거래액", (2026, "1월 증감")])
    assert v[0] in "+△" and v.endswith("%"), f"증감 서식이 깨졌어요 — {v}"
    # 비율 지표는 %p 차이
    r = str(disp.loc["가입율", (2026, "1월 증감")])
    assert r.endswith("%p"), f"비율 지표는 %p여야 해요 — {r}"
    # 값 칸은 평소대로 포맷된다
    assert "원" in str(disp.loc["첫구매 거래액", (2026, "1월")])
    # 색도 입어야 한다(△ 빨강 / + 초록)
    ctx = sty._compute().ctx
    assert any(v for v in ctx.values()), "증감 칸에 색이 안 입었어요"


@case
def t_trend_delta_survives_missing_prior_year():
    """전년이 없는 달은 '–'로 두고 표는 그대로 뜬다."""
    df = synth_store()
    df = df[df["year"] != 2025]                    # 전년을 통째로 없앤다
    tbl = W.trend_table(df, "월", W.METRICS7, [2026], delta_year=2026)
    assert not tbl.empty, "전년이 없다고 표가 통째로 비면 안 돼요"
    dcols = [c for c in tbl.columns if "증감" in c[1]]
    assert dcols, "증감 칸이 아예 없어요"
    assert set(tbl.loc["첫구매 거래액", dcols]) == {"–"}, \
        tbl.loc["첫구매 거래액", dcols].tolist()


@case
def t_month_trim_drops_only_the_leading_zero():
    """'08월 2주차' → '8월 2주차'. 10~12월은 그대로 둔다(0 앞이 숫자라 안 걸린다)."""
    for src, want in (("08월 2주차", "8월 2주차"), ("01월 5주차", "1월 5주차"),
                      ("09월", "9월"), ("10월 1주차", "10월 1주차"),
                      ("11월 3주차", "11월 3주차"), ("12월", "12월"),
                      ("08월 4주차 증감", "8월 4주차 증감"),   # 증감 칸도 같이 다듬긴다
                      # 달이 문자열 **가운데**에 와도 걸려야 한다 — 지금은 늘 맨 앞이라
                      # lstrip("0")로도 통과하지만, 앞에 연도가 붙는 순간 조용히 실패한다
                      ("2026년 08월 2주차", "2026년 8월 2주차"),
                      ("3/1", "3/1"), ("2026. 08", "2026. 08"),
                      ("", ""), (None, None)):
        assert W.month_trim(src) == want, f"{src!r} → {W.month_trim(src)!r} (기대 {want!r})"
    assert W.week_disp(2026, "08월 2주차") == "2026년 8월 2주차"
    assert W.week_disp(2026, "10월 1주차") == "2026년 10월 1주차"
    assert W.week_disp(2026, None) == "-"


@case
def t_stored_labels_keep_their_zero():
    """다듬는 건 **표시뿐**이다 — 저장 라벨이 바뀌면 조회 키가 어긋나 조인이 깨진다."""
    df = synth_store()
    labs = set(df[df["gran"] == "주"]["label"])
    assert any(l.startswith("0") for l in labs), f"합성본에 0으로 시작하는 주차가 없어요 — {sorted(labs)[:5]}"
    # detail_periods가 내는 라벨도 zero-pad를 유지해야 한다
    import pandas as _pd
    pr = W.detail_periods(_pd.Series(["20260810", "20260811"]), "주")
    got = [str(v) for v in pr["label"]]
    assert got and all(v.startswith("08월") for v in got), got
    # 화면 헬퍼를 거치면 그때 다듬긴다
    assert W.month_trim(got[0]).startswith("8월"), got[0]


@case
def t_trend_columns_show_trimmed_labels():
    """추이표 열 머리도 다듬어서 나온다 — 조회는 원래 라벨로 이미 끝났다."""
    df = synth_store()
    tbl = W.trend_table(df, "주", W.METRICS7, [2026])
    labs = [c[1] for c in tbl.columns]
    assert labs, "주차 열이 없어요"
    assert not [l for l in labs if l.startswith("0")], f"앞자리 0이 남았어요 — {labs[:6]}"
    # 값은 그대로 나와야 한다(라벨을 바꿔서 조회가 깨지지 않았는지)
    assert tbl.notna().any().any(), "라벨을 다듬다 조회가 깨졌어요"


def pushall_xlsx(days=400, end=datetime.date(2026, 9, 8), tot=1500, yes=650,
                 rows=("Push Total", "Push Y Cnt", "Push N Cnt")):
    """앱 보유 무관 수신동의 원천 — 실파일과 같은 모양의 xlsx를 만든다.

    1행 날짜(**연도 없음**) · 2행 요일 · 3~5행 Total/Y/N. `rows`를 줄이면
    '세 줄이 다 있어야 잡는다'를 검사할 수 있다.
    """
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Sheet 1"
    ds = [end - datetime.timedelta(days=i) for i in range(days)][::-1]
    for j, d in enumerate(ds, start=2):
        ws.cell(1, j, f"{d.month}/{d.day}")
        ws.cell(2, j, "(월)")
    vals = {"Push Total": tot, "Push Y Cnt": yes, "Push N Cnt": tot - yes}
    for i, name in enumerate(rows, start=3):
        ws.cell(i, 1, name)
        for j in range(2, len(ds) + 2):
            ws.cell(i, j, vals[name])
    bio = io.BytesIO(); wb.save(bio)
    return bio.getvalue()


@case
def t_pushall_source_is_parsed():
    """Total/Y/N 세 줄짜리 원천을 읽어 동의·대상 두 지표로 쌓는다."""
    d = W.parse_pushall_file("Push_14.xlsx", pushall_xlsx(days=400))
    assert not d.empty, "못 읽었어요"
    assert set(d["metric"]) == {"앱푸시동의_앱무관", "앱푸시대상_앱무관"}, set(d["metric"])
    # Push N Cnt는 Total-Y로 언제든 나오는 파생이라 안 쌓는다
    assert not [m for m in set(d["metric"]) if "미동의" in m or "N" in m]
    assert set(d["gran"]) == {"일"}, set(d["gran"])
    # 연도는 상대값으로 나온다 — 절대연도는 combine_files가 마스터와 맞춰 정한다
    assert set(d["year"]) == {0, -1}, sorted(set(d["year"]))
    assert d["label"].str.contains("/").all()


@case
def t_pushall_needs_all_three_rows():
    """세 줄이 다 있어야 이 원천이다 — 한 줄만 보고 잡으면 기존 PUSH 표를 가로챈다."""
    for rows in (("Push Total",), ("Push Total", "Push Y Cnt")):
        d = W.parse_pushall_file("Push_14.xlsx", pushall_xlsx(days=40, rows=rows))
        assert d.empty, f"{rows} 만으로 잡았어요 — {set(d['metric']) if len(d) else ''}"


@case
def t_pushall_does_not_hijack_the_old_push_file():
    """기존 PUSH 원천은 그대로 기존 파서가 가져가야 한다."""
    import inspect as _i
    src = pathlib.Path(APP).read_text(encoding="utf-8")
    i = src.index("pa = parse_pushall_file(n, b)")
    j = src.index("pf = parse_push_file(n, b)")
    assert i < j, "새 파서가 기존 PUSH 파서 뒤에 있어요 — 파일명이 PUSH라 먼저 가로채집니다"
    # 기존 PUSH 표(섹션×행 구조)엔 Total/Y/N 세 줄이 없으니 새 파서가 안 잡는다
    assert "{\"pushtotal\", \"pushycnt\", \"pushncnt\"} <= have" in src


@case
def t_pushall_years_resolve_against_the_master():
    """연도가 없는 원천이라 마스터 일자와 겹침으로 절대연도를 정한다."""
    b = pushall_xlsx(days=400)
    comb = W.combine_files((("Push_14.xlsx", b),))
    pa = comb[comb["metric"] == "앱푸시동의_앱무관"]
    assert not pa.empty, "결합 후 사라졌어요"
    yrs = sorted(set(pa["year"]))
    assert len(yrs) == 2 and yrs[1] - yrs[0] == 1, yrs
    assert yrs[1] >= 2026, f"최근 해가 이상해요 — {yrs}"
    # sortkey가 붙어야 기간 정렬이 된다
    assert (pa["sortkey"] > 0).all(), "sortkey가 안 붙었어요"


@case
def t_pushall_rolls_up_to_the_period_as_a_daily_mean():
    """일별로만 쌓이는 원천이라 주·월 칸은 일평균으로 묶는다."""
    b = pushall_xlsx(days=400, tot=1500, yes=650)
    comb = W.combine_files((("Push_14.xlsx", b),))
    y = max(comb[comb["metric"] == "앱푸시동의_앱무관"]["year"])
    v, approx, n = W.push_period_avg(comb, y, "8월", metric="앱푸시동의_앱무관")
    assert n == 31, f"8월인데 {n}일치예요"
    assert abs(v - 650) < 1e-6, v
    assert not approx, "달은 라벨에서 바로 읽으니 근사가 아니에요"
    # 주는 원천에 주차가 없어 근사다
    _, approx_w, _ = W.push_period_avg(comb, y, "08월 2주차", metric="앱푸시동의_앱무관")
    assert approx_w, "주차는 근사라고 알려야 해요"


@case
def t_funnel_shows_both_consent_denominators():
    """⑤에 두 줄이 나란히 — 앱 보유 무관 / 앱 포함."""
    store = pd.concat([synth_store(), synth_appinstall_store(3000)], ignore_index=True)
    extra = W.combine_files((("Push_14.xlsx", pushall_xlsx(days=400)),))
    at = _open(store=pd.concat([store, extra], ignore_index=True),
               mode="월누적(MTD) — 전년 동월")
    fr = [f for f in _frames(at) if f.index.name == "비율"]
    assert fr, f"비율 표가 없어요 — {[f.index.name for f in _frames(at)]}"
    idx = list(fr[0].index)
    assert "앱푸시(앱보유무관)/신규회원" in idx, idx
    assert "앱푸시(앱포함)/신규회원" in idx, idx
    # 앱 보유 무관이 먼저 온다(분모가 더 넓은 쪽부터 읽는다)
    assert idx.index("앱푸시(앱보유무관)/신규회원") < idx.index("앱푸시(앱포함)/신규회원"), idx
    v = str(fr[0].loc["앱푸시(앱보유무관)/신규회원", "2026년"])
    assert v.endswith("%") and v != "–", v
    assert any("앱을 안 깔아서 못 받는 몫" in t for t in _texts(at)), \
        "두 줄 차이가 뭔지 안 밝혔어요"


@case
def t_missing_pushall_source_says_why():
    """원천이 없으면 빈 줄을 남기지 말고 왜 없는지 말한다."""
    store = pd.concat([synth_store(), synth_appinstall_store(3000)], ignore_index=True)
    at = _open(store=store, mode="월누적(MTD) — 전년 동월")
    fr = [f for f in _frames(at) if f.index.name == "비율"]
    assert fr, "비율 표가 없어요"
    idx = list(fr[0].index)
    assert "앱푸시(앱보유무관)/신규회원" not in idx, "원천이 없는데 빈 줄이 남았어요"
    # **예전 화면 그대로여야 한다** — 구분할 대상이 하나뿐인데 `(앱포함)`을 달아 두면
    # 있지도 않은 다른 지표를 찾게 된다
    assert "신규회원 앱 수신동의율" in idx, idx
    assert not [i for i in idx if "앱포함" in str(i)], idx
    # 안 올린 데이터를 채근하는 문구도 안 띄운다(상시로 뜨면 그것도 소음이다)
    assert not [t for t in _texts(at) if "원천이 없어요" in t], "채근 문구가 떠 있어요"


@case
def t_trend_carries_both_consent_rates_not_install_rate():
    """추이표의 비율 칸은 **동의율 둘**뿐이다 — 설치율까지 섞으면 어느 이야기인지 헷갈린다."""
    store = pd.concat([synth_store(), synth_appinstall_store(3000)], ignore_index=True)
    extra = W.combine_files((("Push_14.xlsx", pushall_xlsx(days=400)),))
    at = _open(store=pd.concat([store, extra], ignore_index=True),
               mode="월누적(MTD) — 전년 동월")
    fr = [f for f in _frames(at) if f.index.name == "기간" and "앱 신규설치" in f.columns]
    assert fr, "추이표가 없어요"
    cols = list(fr[0].columns)
    assert "가입자 대비 설치율" not in cols, f"설치율이 남았어요 — {cols}"
    for c in ("앱푸시(앱보유무관)/신규회원", "앱푸시(앱포함)/신규회원"):
        assert c in cols, f"«{c}» 칸이 없어요 — {cols}"
    # 분모가 넓은 쪽이 먼저 온다 (위 비율 표와 같은 순서)
    assert cols.index("앱푸시(앱보유무관)/신규회원") < cols.index("앱푸시(앱포함)/신규회원"), cols
    v = str(fr[0]["앱푸시(앱보유무관)/신규회원"].iloc[0])
    assert v.endswith("%") and v != "–", v
    assert any("앱을 안 깔아서 못 받는 몫" in t for t in _texts(at)), \
        "두 분모가 다르다는 걸 안 밝혔어요"


@case
def t_trend_drops_pushall_columns_when_source_is_missing():
    """원천이 없으면 '–'만 늘어선 칼럼을 남기지 말고 통째로 뺀다."""
    store = pd.concat([synth_store(), synth_appinstall_store(3000)], ignore_index=True)
    at = _open(store=store, mode="월누적(MTD) — 전년 동월")
    fr = [f for f in _frames(at) if f.index.name == "기간" and "앱 신규설치" in f.columns]
    assert fr, "추이표가 없어요"
    cols = list(fr[0].columns)
    assert not [c for c in cols if "앱보유무관" in c], f"빈 칼럼이 남았어요 — {cols}"
    # 이름도 예전 그대로 — 「앱포함」은 견줄 대상이 있을 때만 붙인다
    assert not [c for c in cols if "앱포함" in c], cols
    assert "앱푸시 수신동의" in cols and "신규회원 수신동의율" in cols, cols
    # 설치율은 원천과 무관하게 뺀 채로 둔다(따로 요청받은 정리다)
    assert "가입자 대비 설치율" not in cols, cols


@case
def t_channel_decomposition_opens_on_revenue():
    """③은 **첫구매 거래액**으로 열린다 — 목록 순서대로 두면 비회원트래픽이 잡힌다."""
    at = _open()
    sel = _sel_step(at)
    assert sel.value == "첫구매 거래액", f"기본값이 «{sel.value}» 예요 — {list(sel.options)}"
    # 첫 칸이 아니라는 걸 못 박아 둔다(순서를 바꿔도 기본값은 거래액이어야 한다)
    assert list(sel.options)[0] != "첫구매 거래액", \
        "목록 첫 칸이 거래액이면 이 검사가 아무것도 안 잡아요"
    # 사용자가 고른 값은 안 덮는다
    sel.set_value("가입자수"); at.run()
    assert not at.exception, at.exception[0].value
    assert _sel_step(at).value == "가입자수", "고른 값이 기본값으로 되돌아갔어요"


@case
def t_guard_select_default_only_seeds_once():
    """default는 옵션에 있을 때만, 세션이 비었을 때만 심는다."""
    import streamlit as _st
    _st.session_state.clear()
    W.guard_select("k1", ["a", "b"], default="b")
    assert _st.session_state["k1"] == "b"
    W.guard_select("k1", ["a", "b"], default="a")          # 이미 있으면 안 덮는다
    assert _st.session_state["k1"] == "b"
    W.guard_select("k2", ["a", "b"], default="z")          # 옵션 밖이면 안 심는다
    assert "k2" not in _st.session_state
    _st.session_state["k3"] = "사라진값"                     # 옵션 밖 옛 값은 치우고
    W.guard_select("k3", ["a", "b"], default="a")          # 그 자리에 기본값
    assert _st.session_state["k3"] == "a"
    for k in ("k1", "k2", "k3"):
        _st.session_state.pop(k, None)


@case
def t_app_rates_say_they_are_daily_means():
    """비율을 '가입한 사람의 몇 %가 앱을 깔았다'로 읽지 않게 못 박는다.

    원천이 가입↔설치를 개인 단위로 잇지 않아 거기까진 알 수 없다. 안 적어 두면
    두 일평균을 나눈 값을 개인 추적 결과로 읽는다.
    """
    store = pd.concat([synth_store(), synth_appinstall_store(3000)], ignore_index=True)
    at = _open(store=store, mode="월누적(MTD) — 전년 동월")
    txt = _texts(at)
    assert any("일평균끼리 나눈 값" in t for t in txt), "일평균 비율이라는 안내가 없어요"
    assert any("개인 단위로" in t for t in txt), "개인 추적이 아니라는 안내가 없어요"
    # 신규 설치 기준(재설치 제외)이라는 안내는 카드 아래에 이미 있다
    assert any("재설치는 빼고" in t for t in txt), "신규 설치 기준 안내가 없어요"


@case
def t_picked_row_reads_cell_selection():
    """_picked_row는 셀 선택·행 선택 둘 다 받고, 범위를 벗어나면 None."""
    class _Ev:
        def __init__(self, sel): self.selection = sel
    assert W._picked_row(_Ev({"cells": [(2, "2026년")], "rows": []}), 5) == 2
    assert W._picked_row(_Ev({"cells": [], "rows": [1]}), 5) == 1      # 옛 세션 호환
    assert W._picked_row(_Ev({"cells": [(9, "x")], "rows": []}), 5) is None
    assert W._picked_row(_Ev({"cells": [], "rows": []}), 5) is None
    assert W._picked_row(_Ev(None), 5) is None
    assert W._picked_row({"selection": {"cells": [[0, "a"]]}}, 3) == 0  # dict로 오는 경우


@case
def t_app_detail_swaps_prior_year_for_shares():
    """앱설치 상세 — 전년이 한 칸도 없으면 그 열 대신 전체설치 대비 비중을 낸다."""
    store = pd.concat([synth_store(), synth_appinstall_store(3000)], ignore_index=True)
    at = _open(store=store, mode="월누적(MTD) — 전년 동월")
    fr = [f for f in _frames(at) if f.index.name == "지표" and "전체설치" in f.index]
    assert fr, f"앱설치 상세 표가 없어요 — {[list(f.index) for f in _frames(at)]}"
    t = fr[0]
    assert "전체설치 대비" in t.columns, list(t.columns)
    assert not [c for c in t.columns if "전년" in str(c)], \
        f"전년 칸이 통째로 비는데 남아 있어요 — {list(t.columns)}"
    assert "신규설치" in t.index, "전체설치 = 신규 + 재설치가 한 표에서 닫혀야 해요"
    # 재설치 비중은 재설치 ÷ 전체설치 — 100%를 넘을 수 없다
    _rr = str(t.loc["재설치", "전체설치 대비"])
    assert _rr.endswith("%") and 0 < float(_rr.rstrip("%")) < 100, _rr
    # Push 활성 기기는 잔고라 비중이 없다
    assert t.loc["Push활성기기", "전체설치 대비"] == "–"
    assert any("전체설치 대비 비중을 넣었어요" in x for x in _texts(at)), "왜 바꿨는지 안 밝혔어요"


@case
def t_channel_decomposition_accepts_ratio_metrics():
    """당일가입CR 같은 비율도 고를 수 있어야 한다 — 다만 워터폴이 아니라 채널별 변화다."""
    at = _open()
    opts = list(_sel_step(at).options)
    for m in ("당일가입CR", "가입율", "첫구매 객단가"):
        assert m in opts, f"«{m}»를 못 골라요 — {opts}"
    _sel_step(at).set_value("당일가입CR"); at.run()
    assert not at.exception, at.exception[0].value
    assert any("채널을 더해도 전체가 안 돼요" in t for t in _texts(at)), \
        "비율 지표 주의 문구가 없어요"


@case
def t_app_block_states_its_period():
    """⑤는 한참 내려온 자리라 어느 기간 값인지 다시 말해야 한다."""
    store = pd.concat([synth_store(), synth_appinstall_store(3000)], ignore_index=True)
    at = _open(store=store)
    assert any(t.startswith("기준:") and "일평균" in t for t in _texts(at)), \
        [t for t in _texts(at) if "기준" in t]


@case
def t_app_block_shows_recent_periods():
    """카드 한 장으로는 '이번이 낮은 건지 원래 그런 건지'를 못 본다."""
    store = pd.concat([synth_store(), synth_appinstall_store(3000)], ignore_index=True)
    at = _open(store=store)
    fr = [f for f in _frames(at) if f.index.name == "기간" and "앱 신규설치" in f.columns]
    assert fr, f"최근 추이 표가 없어요 — {[f.index.name for f in _frames(at)]}"
    assert len(fr[0]) >= W.APP_TREND_N["주"], f"주 단위는 8개 이상 — {len(fr[0])}"
    # 앱 원천이 마스터보다 짧게 끝나도 **지금 보고 있는 기간**은 빈 줄로라도 서야 한다 —
    # 표에서 통째로 빠지면 '왜 카드가 비었지'가 안 풀린다
    assert sum("◀" in str(i) for i in fr[0].index) == 1, list(fr[0].index)
    # 단위는 위 비교 기준을 따라간다
    at2 = _open(store=store, mode="월누적(MTD) — 전년 동월")
    fr2 = [f for f in _frames(at2) if f.index.name == "기간" and "앱 신규설치" in f.columns]
    assert fr2 and len(fr2[0]) >= W.APP_TREND_N["월"], \
        f"월 단위는 6개 이상 — {len(fr2[0]) if fr2 else None}"
    assert sum("◀" in str(i) for i in fr2[0].index) == 1, list(fr2[0].index)


@case
def t_app_block_reports_coverage():
    """기간은 맞는데 원천이 아직 안 닿았을 수 있다 — 어디까지 들어왔는지 말한다."""
    store = pd.concat([synth_store(), synth_appinstall_store(3000)], ignore_index=True)
    at = _open(store=store)
    assert any("앱설치 원천은" in t and "까지 들어와 있어요" in t for t in _texts(at)), \
        "커버리지 안내가 없어요"


@case
def t_empty_app_period_says_why():
    """마스터는 09월 4주차까지인데 앱 원천은 09월 1주차까지다 — 그때 빈 카드만 두면 안 된다."""
    store = pd.concat([synth_store(), synth_appinstall_store(3000)], ignore_index=True)
    at = _open(store=store)
    warn = " ".join(str(w.value) for w in at.warning)
    assert "앱 데이터가 없어요" in warn and "까지 들어와 있어요" in warn, warn


@case
def t_factor_split_shows_where_it_leaked():
    """`거래액 = 상품UV × 상품CR × 객단가` — 어느 요인이 끌어내렸는지 짚어야 한다."""
    at = _open()
    fr = [f for f in _frames(at) if f.index.name == "요인"]
    assert fr, f"요인 분해 표가 없어요 — {[f.index.name for f in _frames(at)]}"
    idx = " ".join(str(i) for i in fr[0].index)
    for nick in ("유입", "전환", "객단가"):
        assert nick in idx, (nick, list(fr[0].index))
    # 제목은 접이식 라벨로 옮겼다 — _texts는 expander 라벨을 안 모은다
    assert any("어디에서 빠졌나" in str(e.label) for e in at.expander), \
        f"요인 분해 제목이 없어요 — {[str(e.label) for e in at.expander]}"


@case
def t_factor_split_refuses_when_it_cannot():
    """0·결측이 섞이면 로그가 정의되지 않는다 — 숫자를 지어내면 안 된다."""
    assert W.factor_split([1.0, 0.0, 1.0], [1.0, 1.0, 1.0]) is None
    assert W.factor_split([1.0, 1.0], [1.0, np.nan]) is None
    got = W.factor_split([10.0, 2.0], [20.0, 2.0])
    assert got is not None
    parts, total = got
    assert abs(sum(parts) - total) < 1e-6, (parts, total)   # 합이 실제 증감과 같다


@case
def t_rollup_covers_every_selectable_metric():
    """④에서 고를 수 있는 지표는 **전부** 조직 합산 표가 나와야 한다.

    예전엔 상품UV·상품CR을 막아 뒀는데, 원천이 조직×카테고리로 주는 값이라 UV는
    더하면 되고 CR은 항등식으로 되만들 수 있다. 막아 두면 '가방이 어느 조직에서든
    전환이 빠졌나'를 조직 하나씩 들어가 봐야만 알 수 있다.
    막는 분기는 남겨 두되(고를 수 없는 지표가 생길 때를 위해) 지금은 아무도 안 밟는다.
    """
    at = _open()
    opts = list(_sel_oc(at).options)
    assert set(opts) <= set(W.FUNNEL_ROLLUP_METS), \
        f"모을 수 없는 지표를 고를 수 있어요 — {set(opts) - set(W.FUNNEL_ROLLUP_METS)}"
    for met in opts:
        _sel_oc(at).set_value(met); at.run()
        assert not at.exception, f"{met}: {at.exception[0].value if at.exception else ''}"
        assert [f for f in _frames(at) if f.index.name == "카테고리"], \
            f"«{met}» 합산 표가 없어요"


def main():
    fails, cwd = [], os.getcwd()
    for fn in CASES:
        try:
            fn()
            print(f"  OK   {fn.__name__}")
        except Exception as e:                            # noqa: BLE001
            print(f"  FAIL {fn.__name__}: {str(e)[:300]}")
            fails.append(fn.__name__)
        finally:
            os.chdir(cwd)              # _open이 임시 폴더에 머무르므로 여기서 되돌린다
    print()
    if fails:
        print(f"실패 {len(fails)}건: {fails}")
        return 1
    print(f"첫구매 퍼널 페이지 테스트 {len(CASES)}건 통과 ✅")
    return 0


if __name__ == "__main__":
    sys.exit(main())
